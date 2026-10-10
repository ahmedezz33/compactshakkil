# -*- coding: utf-8 -*-
"""
ocr_export.py — تصديرُ ناتجِ OCR إلى ملفّ Word (.docx) عربيٍّ صحيحِ الاتّجاه (RTL).

- يحافظُ على الأقسام (العنوان/المتن/الحواشي) إن كان النصُّ مهيكَلًا.
- اتّجاهٌ من اليمين، خطٌّ عربيّ، عناوينُ الأقسامِ بارزة.
- اختياريًّا يستردُّ تشكيلَ الآياتِ القرآنيّةِ من المصحفِ الموثوق (restore_quran_tashkeel).

يتطلّبُ: pip install python-docx
"""
import os, re

_LBL = re.compile(r"(?:▓▓▓|=)\s*(العنوان|المتن|الحواشي)\s*(?:▓▓▓|=)")

def _sections(text):
    """يُرجِعُ [(عنوانُ القسم أو None, نصّه)] بالترتيب. إن لم تُوجَدْ علاماتٌ: المتنُ كلُّه."""
    if not text: return []
    if not _LBL.search(text):
        return [(None, text.strip())]
    out=[]; pos=0; cur=None; last=0
    for m in _LBL.finditer(text):
        if cur is not None or last<m.start():
            seg=text[last:m.start()].strip()
            if seg: out.append((cur, seg))
        cur=m.group(1); last=m.end()
    seg=text[last:].strip()
    if seg: out.append((cur, seg))
    return out

def _set_rtl(paragraph):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    pPr=paragraph._p.get_or_add_pPr()
    b=OxmlElement("w:bidi"); pPr.append(b)
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for run in paragraph.runs:
        rPr=run._element.get_or_add_rPr()
        rtl=OxmlElement("w:rtl"); rtl.set(qn("w:val"),"1"); rPr.append(rtl)

def _font(run, name, size, bold=False):
    from docx.oxml.ns import qn
    from docx.shared import Pt
    run.font.size=Pt(size); run.bold=bold
    run.font.name=name
    rPr=run._element.get_or_add_rPr()
    rf=rPr.find(qn("w:rFonts"))
    if rf is None:
        from docx.oxml import OxmlElement
        rf=OxmlElement("w:rFonts"); rPr.append(rf)
    rf.set(qn("w:cs"), name); rf.set(qn("w:ascii"), name); rf.set(qn("w:hAnsi"), name)

def export_docx(text, out_path, structured=True, quran="strong",
                font="Arial", body_size=14, title="نصّ مُستخرَج (OCR)"):
    """quran: "off" | "braced" (المقوّسةُ فقط، آمن) | "strong" (المقوّسةُ وغيرُها)."""
    try:
        import docx
        from docx import Document
    except Exception:
        raise RuntimeError("python-docx غير مثبَّت. ثبّتْه: pip install python-docx")
    if quran and quran!="off":
        try:
            import quran_audio
            if quran=="strong" and hasattr(quran_audio,"restore_quran_strong"):
                text=quran_audio.restore_quran_strong(text)   # مقوّسةٌ وغيرُها ← المصحف
            else:
                text=quran_audio.restore_quran_tashkeel(text) # المقوّسةُ فقط (آمن)
        except Exception as e:
            print("quran restore skip:", e)

    doc=Document()
    # اتّجاهُ المقطعِ RTL
    try:
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        sect=doc.sections[0]._sectPr
        b=OxmlElement("w:bidi"); sect.append(b)
    except Exception: pass
    # النمطُ الافتراضيّ عربيّ
    try:
        from docx.shared import Pt
        st=doc.styles["Normal"]; st.font.name=font; st.font.size=Pt(body_size)
    except Exception: pass

    secs=_sections(text) if structured else [(None, text)]
    HEAD={"العنوان":"العنوان / الترويسة","المتن":"المتن","الحواشي":"الحواشي"}
    first=True
    for label, body in secs:
        if label and label in HEAD:
            h=doc.add_paragraph(); r=h.add_run(HEAD[label]); _font(r, font, body_size+4, bold=True); _set_rtl(h)
        for line in body.split("\n"):
            line=line.strip()
            if not line: continue
            p=doc.add_paragraph(); r=p.add_run(line); _font(r, font, body_size); _set_rtl(p)
        first=False
    doc.save(out_path)
    return out_path

if __name__=="__main__":
    import sys
    inp=sys.argv[1]; out=sys.argv[2] if len(sys.argv)>2 else "ocr_out.docx"
    txt=open(inp,encoding="utf-8").read()
    print("✅", export_docx(txt, out))
