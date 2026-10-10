"""
build_jawf_review.py — بناء ملفّات مراجعة Jawf-DB.

المخرجات:
    jawf_review/batch_01_top.xlsx    → أعلى ~1000 جذر استخدامًا
    jawf_review/batch_02.xlsx        → الجذور 1001-2000
    jawf_review/batch_03.xlsx        → الجذور 2001-3000
    jawf_review/batch_04.xlsx        → الجذور 3001-4000
    jawf_review/batch_05.xlsx        → الجذور 4001-5000
    jawf_review/batch_06_rare.xlsx   → الجذور النادرة 5001+

الاستخدام:
    python lexicon/build_jawf_review.py
"""

import sqlite3
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DB_PATH = 'lexicon/arabic_lexicon.db'
OUTPUT_DIR = Path('jawf_review')
BATCH_SIZE = 1000
NUM_BATCHES = 6


COLUMNS = [
    ("id", 6, "auto"),
    ("root_camel", 12, "auto"),
    ("derivatives_count", 12, "auto"),
    ("weak_type", 15, "auto"),
    ("sample_words", 45, "auto"),
    ("root_correct", 12, "manual"),
    ("weak_letter", 10, "manual"),
    ("meaning_primary", 30, "manual"),
    ("meaning_secondary", 30, "manual"),
    ("has_ambiguity", 12, "manual"),
    ("alternative_root", 14, "manual"),
    ("lisan_arab_ref", 14, "manual"),
    ("wasit_ref", 12, "manual"),
    ("taj_arus_ref", 14, "manual"),
    ("qamus_muhit_ref", 14, "manual"),
    ("buckwalter_ref", 14, "manual"),
    ("verified_by", 12, "manual"),
    ("verified_date", 12, "manual"),
    ("notes", 40, "manual"),
]


def classify_weak_type(root):
    parts = root.split('.')
    if len(parts) != 3:
        return "غير منتظم"
    positions = [i for i, p in enumerate(parts) if p == '#']
    if len(positions) == 1:
        return {0: "مثال", 1: "أجوف", 2: "ناقص"}[positions[0]]
    elif len(positions) == 2:
        if 0 in positions and 2 in positions:
            return "لفيف مفروق"
        return "لفيف مقرون"
    return "غامض"


def get_all_weak_roots():
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()

    print("📖 استخراج الجذور المعتلّة من المعجم...")
    c.execute("""
        SELECT root, COUNT(*) as derivative_count
        FROM roots
        WHERE root LIKE '%#%'
        GROUP BY root
        ORDER BY derivative_count DESC
    """)

    roots_data = []
    all_rows = c.fetchall()
    print(f"   وُجد: {len(all_rows):,} جذر معتلّ")

    print("📝 جلب عيّنات لكل جذر...")
    for root, count in all_rows:
        c.execute("""
            SELECT DISTINCT w.word
            FROM words w
            JOIN roots r ON w.id = r.word_id
            WHERE r.root = ?
            ORDER BY w.frequency DESC
            LIMIT 5
        """, (root,))
        samples = [row[0] for row in c.fetchall()]

        roots_data.append({
            "root": root,
            "count": count,
            "samples": " | ".join(samples),
            "weak_type": classify_weak_type(root),
        })

    conn.close()
    return roots_data


def create_excel_batch(batch_num, batch_data, total_batches, is_last):
    OUTPUT_DIR.mkdir(exist_ok=True)

    if batch_num == 1:
        filename = f"batch_{batch_num:02d}_top.xlsx"
    elif is_last:
        filename = f"batch_{batch_num:02d}_rare.xlsx"
    else:
        filename = f"batch_{batch_num:02d}.xlsx"

    filepath = OUTPUT_DIR / filename

    wb = Workbook()
    ws = wb.active
    ws.title = f"Batch {batch_num}"
    ws.sheet_view.rightToLeft = True

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill_auto = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_fill_manual = PatternFill(start_color="C55A11", end_color="C55A11", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right_align = Alignment(horizontal="right", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC'),
    )

    for col_idx, (col_name, col_width, col_type) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        if col_type == "auto":
            cell.fill = header_fill_auto
        else:
            cell.fill = header_fill_manual

        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    ws.row_dimensions[1].height = 30

    start_id = (batch_num - 1) * BATCH_SIZE + 1
    row_data_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    for row_idx, item in enumerate(batch_data, start=2):
        row_id = start_id + row_idx - 2

        values = [
            row_id,
            item["root"],
            item["count"],
            item["weak_type"],
            item["samples"],
        ] + [""] * 14

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = right_align if isinstance(val, str) else center_align
            cell.border = thin_border
            if col_idx <= 5:
                cell.fill = row_data_fill

        ws.row_dimensions[row_idx].height = 22

    ws.freeze_panes = "A2"

    instructions_ws = wb.create_sheet("التعليمات", 0)
    instructions_ws.sheet_view.rightToLeft = True

    instructions = [
        ["Jawf-DB — ملفّ مراجعة الجذور المعتلّة"],
        [""],
        [f"الدفعة رقم: {batch_num} من {total_batches}"],
        [f"عدد الجذور في هذه الدفعة: {len(batch_data)}"],
        [""],
        ["الأعمدة الملوّنة بالأزرق: تلقائية (لا تعدّلها)"],
        ["الأعمدة الملوّنة بالبرتقالي: للمراجعة اليدوية"],
        [""],
        ["بروتوكول المراجعة لكل جذر:"],
        ["  1. انظر إلى root_camel والمشتقّات في sample_words"],
        ["  2. خمّن الحرف المعتلّ (و/ي/ا) بناءً على المشتقّات"],
        ["  3. افتح المكتبة الشاملة وابحث عن الجذر"],
        ["  4. سجّل مرجع لسان العرب (ج ص) في lisan_arab_ref"],
        ["  5. سجّل مرجع القاموس المحيط في qamus_muhit_ref"],
        ["  6. اكتب المعنى الرئيسي في meaning_primary"],
        ["  7. إذا كان هناك جذر آخر بنفس CAMeL — سجّل has_ambiguity=نعم و alternative_root"],
        ["  8. احفظ الملفّ (Ctrl+S) وانتقل للجذر التالي"],
        [""],
        ["مثال:"],
        ["  root_camel: ق.#.ل"],
        ["  root_correct: ق.و.ل"],
        ["  weak_letter: و"],
        ["  meaning_primary: القول والكلام"],
        ["  has_ambiguity: نعم"],
        ["  alternative_root: ق.ي.ل"],
        ["  lisan_arab_ref: ج11 ص572"],
        [""],
        ["نصيحة: ابدأ بأوّل 20 جذرًا لتعتاد على السير، ثمّ اعمل بإيقاع منتظم."],
    ]

    for row_idx, row_data in enumerate(instructions, start=1):
        for col_idx, value in enumerate(row_data, start=1):
            cell = instructions_ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = right_align
            if row_idx == 1:
                cell.font = Font(bold=True, size=14, color="1F4E79")
            elif row_idx in [3, 4, 9, 19]:
                cell.font = Font(bold=True, size=11)

    instructions_ws.column_dimensions['A'].width = 80

    wb.save(filepath)
    file_size = filepath.stat().st_size / 1024
    print(f"   ✓ {filename}: {len(batch_data)} جذر ({file_size:.1f} KB)")

    return filepath


def main():
    print("=" * 60)
    print("Jawf-DB — بناء ملفّات المراجعة")
    print("=" * 60)

    all_roots = get_all_weak_roots()
    total = len(all_roots)
    print(f"\n📊 الإجمالي: {total:,} جذر معتلّ")

    print(f"\n📦 التقسيم إلى {NUM_BATCHES} دفعات:")

    batches = []
    for i in range(NUM_BATCHES):
        start_idx = i * BATCH_SIZE
        end_idx = start_idx + BATCH_SIZE if i < NUM_BATCHES - 1 else total
        batch = all_roots[start_idx:end_idx]
        if batch:
            batches.append(batch)

    for i, batch in enumerate(batches, start=1):
        is_last = (i == len(batches))
        create_excel_batch(i, batch, len(batches), is_last)

    print(f"\n" + "=" * 60)
    print(f"✅ اكتمل بناء ملفّات المراجعة!")
    print(f"=" * 60)
    print(f"\n📁 الموقع: {OUTPUT_DIR.absolute()}")
    print(f"\nالخطوة التالية:")
    print(f"  1. افتح: jawf_review\\batch_01_top.xlsx")
    print(f"  2. اقرأ ورقة 'التعليمات' أوّلًا")
    print(f"  3. ابدأ المراجعة من أعلى استخدامًا")
    print(f"  4. احفظ باستمرار (Ctrl+S)")


if __name__ == "__main__":
    main()
